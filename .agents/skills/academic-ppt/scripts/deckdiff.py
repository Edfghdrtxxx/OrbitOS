"""deckdiff.py A.pptx B.pptx — shape-level diff (text, geometry, fill, shadow, strike, highlight, notes, images), including groups, tables, chart titles, and master text, plus PowerPoint comments added or removed in B.
Slides are paired across inserts, deletes, duplicates, and reorders. Each text change is its own line. A Comment is kept whole, including on an inserted or deleted slide's notes."""
import sys, hashlib, difflib
from collections import Counter
from pptx import Presentation
from lxml import etree
E = 914400
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}

def _frames(sh):
    if sh.has_text_frame:
        return [sh.text_frame]
    if getattr(sh, "has_table", False):
        return [cell.text_frame for cell in sh.table.iter_cells()]
    return []

def _fill(el, root):
    if not hasattr(el, "xpath"):
        return []
    rgb = el.xpath(f"{root}/a:solidFill/a:srgbClr/@val")
    return (rgb or el.xpath(f"{root}/a:solidFill/a:schemeClr/@val"))[:1]

def _chart_text(sh):
    if not getattr(sh, "has_chart", False):
        return ""
    try:
        el = sh.chart._element
    except Exception:
        return ""
    return "\n".join(t for t in el.xpath(".//*[local-name()='t']/text()") if t.strip())

def _flags(r):
    """Bold, italic, underline, font color, highlight. Strike has its own field."""
    rPr = r._r.rPr
    if rPr is None:
        return ()
    flags = []
    if rPr.get("b") in ("1", "true"): flags.append("b")
    if rPr.get("i") in ("1", "true"): flags.append("i")
    if rPr.get("u") not in (None, "none"): flags.append("u")
    cols = rPr.xpath("./*[local-name()='solidFill']/*[local-name()='srgbClr']/@val")
    schemes = rPr.xpath("./*[local-name()='solidFill']/*[local-name()='schemeClr']/@val")
    hls = rPr.xpath("./*[local-name()='highlight']//*[local-name()='srgbClr']/@val")
    if cols: flags.append("color:" + cols[0])
    elif schemes: flags.append("color:" + schemes[0])
    if hls: flags.append("hl:" + hls[0])
    return tuple(flags)

def _crop(sh):
    rects = sh._element.xpath("./p:blipFill/a:srcRect") if sh.shape_type != 6 else []
    if not rects:
        return ()
    return tuple(int(rects[0].get(k) or 0) for k in ("l", "t", "r", "b"))

def _rot(sh):
    xf = sh._element.xpath("./p:spPr/a:xfrm" if sh.shape_type != 6 else "./p:grpSpPr/a:xfrm")
    if not xf:
        return (0, None, None)
    return (int(xf[0].get("rot") or 0), xf[0].get("flipH"), xf[0].get("flipV"))

def _descr(sh):
    tag = {6: "nvGrpSpPr", 13: "nvPicPr"}.get(sh.shape_type, "nvSpPr")
    nodes = sh._element.xpath(f"./p:{tag}/p:cNvPr")
    return (nodes[0].get("descr") or "") if nodes else ""

def _links(sh):
    if sh.shape_type == 6:
        return ()
    out = []
    for h in sh._element.xpath(".//*[local-name()='hlinkClick']"):
        rid = h.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
        if rid and rid in sh.part.rels:
            out.append(sh.part.rels[rid].target_ref)
        elif h.get("action"):
            out.append(h.get("action"))
    return tuple(out)

def sig(sh):
    el = sh._element
    # A group's own properties live on grpSpPr. Descending would fold every child into the group.
    root = "./p:grpSpPr" if sh.shape_type == 6 else ".//p:spPr"
    fill = _fill(el, root)
    shd_path = "./p:grpSpPr/a:effectLst/a:outerShdw" if sh.shape_type == 6 else ".//a:effectLst/a:outerShdw"
    shd = bool(el.xpath(shd_path)) if hasattr(el, "xpath") else False
    img = hashlib.md5(sh.image.blob).hexdigest()[:8] if sh.shape_type == 13 else ""
    frames = _frames(sh)
    txt = "\n".join(t for t in [tf.text for tf in frames] + [_chart_text(sh)] if t)
    runs = [r for tf in frames for p in tf.paragraphs for r in p.runs]
    strike = " | ".join(r.text for r in runs if r._r.rPr is not None and r._r.rPr.get("strike", "noStrike") != "noStrike")
    geo = tuple(round((v or 0) / E, 2) for v in (sh.left, sh.top, sh.width, sh.height))
    styled = tuple((r.text, _flags(r)) for r in runs if r.text)
    return dict(name=sh.name, txt=txt, geo=geo, fill=fill, shadow=shd, img=img, strike=strike, styled=styled,
                crop=_crop(sh), rot=_rot(sh), descr=_descr(sh), link=_links(sh))

def collect(container, out):
    for sh in container:
        out[sh.shape_id] = sig(sh)
        if sh.shape_type == 6:
            collect(sh.shapes, out)

def slide_comments(slide):
    """Legacy and modern PowerPoint comment texts on this slide."""
    out = []
    for r in slide.part.rels.values():
        if not r.is_external and r.reltype.endswith("/comments"):
            for cm in etree.fromstring(r.target_part.blob).xpath("//*[local-name()='cm']"):
                out.append(" ".join(cm.xpath(".//*[local-name()='text' or local-name()='t']/text()")))
    return out

def deck(path):
    p = Presentation(path)
    out = []
    for s in p.slides:
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide and s.notes_slide.notes_text_frame else ""
        found = {}
        collect(s.shapes, found)
        out.append((found, notes, slide_comments(s)))
    return out

def _comment_spans(s):
    """[start, end) covering a whole comment. A ')' inside (Comment: …) does not end it early, and a newline does not end a Comment: block."""
    spans, i = [], 0
    while True:
        j = s.find("Comment:", i)
        if j < 0:
            break
        if j > 0 and s[j - 1] == "(":
            depth, end = 0, None
            for k in range(j - 1, len(s)):
                if s[k] == "(": depth += 1
                elif s[k] == ")":
                    depth -= 1
                    if depth == 0:
                        end = k + 1
                        break
            spans.append((j - 1, len(s) if end is None else end))
        else:
            spans.append((j, len(s)))
        i = j + 8
    return spans

def _grow(s, start, end, spans):
    for c0, c1 in spans:
        if start < c1 and end > c0:
            start, end = min(start, c0), max(end, c1)
    return start, end

def _snap_left(s, start, earliest):
    i = start
    while i > earliest and not s[i - 1].isspace():
        i -= 1
    return i if i == 0 or s[i - 1].isspace() else start

def _snap_right(s, end, latest):
    i = end
    while i < latest and not s[i].isspace():
        i += 1
    return i if i == len(s) or s[i].isspace() else end

def excerpts(a, b, w=40):
    """Each change in a and b, with about w characters of context. A Comment is kept whole."""
    a, b = str(a), str(b)
    if a == b:
        return []
    ops = [op for op in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes() if op[0] != "equal"]
    if not ops:
        return [(a, b)]
    groups = [[ops[0]]]
    for op in ops[1:]:
        gap_a = op[1] - groups[-1][-1][2]
        gap_b = op[3] - groups[-1][-1][4]
        if gap_a <= w and gap_b <= w:
            groups[-1].append(op)
        else:
            groups.append([op])
    ca, cb = _comment_spans(a), _comment_spans(b)
    intervals = []
    for g in groups:
        i1, i2 = _grow(a, g[0][1], g[-1][2], ca)
        j1, j2 = _grow(b, g[0][3], g[-1][4], cb)
        if intervals and i1 <= intervals[-1][1] and j1 <= intervals[-1][3]:
            p = intervals[-1]
            intervals[-1] = (min(p[0], i1), max(p[1], i2), min(p[2], j1), max(p[3], j2))
        else:
            intervals.append((i1, i2, j1, j2))
    out, prev_a, prev_b = [], 0, 0
    for n, (i1, i2, j1, j2) in enumerate(intervals):
        nxt_a = intervals[n + 1][0] if n + 1 < len(intervals) else len(a)
        nxt_b = intervals[n + 1][2] if n + 1 < len(intervals) else len(b)
        avail_l = 0
        while avail_l < i1 - prev_a and avail_l < j1 - prev_b and a[i1 - 1 - avail_l] == b[j1 - 1 - avail_l]:
            avail_l += 1
        avail_r = 0
        while avail_r < nxt_a - i2 and avail_r < nxt_b - j2 and i2 + avail_r < len(a) and j2 + avail_r < len(b) and a[i2 + avail_r] == b[j2 + avail_r]:
            avail_r += 1
        left = _snap_left(a, i1 - min(avail_l, w), i1 - min(avail_l, w + 40))
        right_a = _snap_right(a, i2 + min(avail_r, w), i2 + min(avail_r, w + 40))
        shift_l, shift_r = i1 - left, right_a - i2
        out.append((_cut(a, left, right_a), _cut(b, j1 - shift_l, j2 + shift_r)))
        prev_a, prev_b = i2, j2
    return out

def format_changes(a, b):
    """Formatting changes on text that still matches. A recolored word stays visible when some other word in the shape was rewritten."""
    ca = [(ch, flags) for text, flags in a for ch in text]
    cb = [(ch, flags) for text, flags in b for ch in text]
    ta, tb = "".join(c for c, _ in ca), "".join(c for c, _ in cb)
    spans = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, ta, tb, autojunk=False).get_opcodes():
        if tag != "equal":
            continue
        k = i1
        while k < i2:
            fk = j1 + (k - i1)
            if ca[k][1] != cb[fk][1]:
                s = k
                while k < i2 and ca[k][1] == ca[s][1] and cb[j1 + (k - i1)][1] == cb[j1 + (s - i1)][1]:
                    k += 1
                spans.append((s, k, j1 + (s - i1)))
            else:
                k += 1
    out = []
    for i1, i2, j1 in spans:
        fo, fn = ca[i1][1], cb[j1][1]
        if fo == fn or not ta[i1:i2].strip():
            continue
        old = "{" + ",".join(fo) + "}" if fo else ""
        new = "{" + ",".join(fn) + "}" if fn else ""
        out.append((ta[i1:i2] + old, tb[j1:j1 + (i2 - i1)] + new))
    return out

def _cut(s, start, end):
    """Mark a context edge that is still mid-token, so a trimmed run of letters is not read as a word."""
    start, end = max(0, start), min(len(s), end)
    pre = "…" if start > 0 and not s[start - 1].isspace() else ""
    post = "…" if end < len(s) and not s[end].isspace() else ""
    return pre + s[start:end] + post

def align(A, B):
    """(i, j) slide pairs; unchanged shape sets match in order, then each changed block pairs by overlap."""
    key = lambda s: tuple(sorted((sid, v["name"]) for sid, v in s[0].items()))
    pairs = []
    for tag, i1, i2, j1, j2 in difflib.SequenceMatcher(None, [key(s) for s in A], [key(s) for s in B], autojunk=False).get_opcodes():
        if tag == "equal": pairs += [(i1 + k, j1 + k) for k in range(i2 - i1)]; continue
        free = set(range(j1, j2))
        for i in range(i1, i2):
            ka = set(key(A[i]))
            score = {j: len(ka & set(key(B[j]))) / len(ka | set(key(B[j]))) for j in free}
            j = max(score, key=score.get, default=None)
            if j is not None and score[j] >= 0.5: pairs.append((i, j)); free.discard(j)
            else: pairs.append((i, None))
        pairs += [(None, j) for j in sorted(free)]
    return _repair_same_layout(A, B, _pair_moves(A, B, pairs, key), key)

def _pair_moves(A, B, pairs, key):
    """Pair a deleted slide with an inserted one when they are the same slide moved.
    SequenceMatcher reports a move as a separate delete and insert."""
    unmatched_a = [i for i, j in pairs if j is None]
    unmatched_b = [j for i, j in pairs if i is None]
    scored = []
    for i in unmatched_a:
        ka = set(key(A[i]))
        for j in unmatched_b:
            kb = set(key(B[j]))
            union = ka | kb
            scored.append((len(ka & kb) / len(union) if union else 0, i, j))
    used_a, used_b, found = set(), set(), []
    for score, i, j in sorted(scored, reverse=True):
        if score < 0.5 or i in used_a or j in used_b:
            continue
        found.append((i, j))
        used_a.add(i)
        used_b.add(j)
    kept = [(i, j) for i, j in pairs if i not in used_a and j not in used_b]
    return sorted(kept + found, key=lambda p: (p[1] if p[1] is not None else p[0] + 0.5))

def _repair_same_layout(A, B, pairs, key):
    """Slides that share a shape layout are interchangeable to the structural pass.
    Re-pair all of them, including a deleted or inserted one, by text."""
    def payload(s):
        body = tuple(sorted((sid, v["txt"], v["strike"], v["styled"]) for sid, v in s[0].items()))
        return (s[1], body)
    a_of, b_of = {}, {}
    for i, s in enumerate(A):
        a_of.setdefault(key(s), []).append(i)
    for j, s in enumerate(B):
        b_of.setdefault(key(s), []).append(j)
    ambiguous = {k for k in set(a_of) | set(b_of) if len(a_of.get(k, [])) > 1 or len(b_of.get(k, [])) > 1}
    if not ambiguous:
        return pairs
    amb_a = {i for k in ambiguous for i in a_of[k]}
    amb_b = {j for k in ambiguous for j in b_of[k]}
    pref = {(i, j) for i, j in pairs if i in amb_a and j in amb_b}
    kept = [(i, j) for i, j in pairs if i not in amb_a and j not in amb_b]
    for k in ambiguous:
        ais, bjs = a_of.get(k, []), b_of.get(k, [])
        buckets = {}
        for j in bjs:
            buckets.setdefault(payload(B[j]), []).append(j)
        used, chosen = set(), {}
        for ai in ais:
            hits = buckets.get(payload(A[ai]), [])
            orig = next((j for i, j in pref if i == ai), None)
            pick = orig if orig in hits and orig not in used else next((h for h in hits if h not in used), None)
            if pick is not None:
                used.add(pick)
                chosen[ai] = pick
        left_a = [ai for ai in ais if ai not in chosen]
        left_b = [j for j in bjs if j not in used]
        for ai, j in zip(left_a, left_b):
            chosen[ai] = j
        kept += [(ai, j) for ai, j in chosen.items()]
        kept += [(ai, None) for ai in left_a[len(left_b):]]
        kept += [(None, j) for j in left_b[len(left_a):]]
    return sorted(kept, key=lambda p: (p[1] if p[1] is not None else p[0] + 0.5))

def chrome(path):
    """Text on slide masters and layouts. A comment typed there never appears in a slide's shape tree."""
    p = Presentation(path)
    items = []
    for mi, m in enumerate(p.slide_masters, 1):
        found = {}
        collect(m.shapes, found)
        items.append((f"master {mi}", found))
        for li, lay in enumerate(m.slide_layouts, 1):
            found = {}
            collect(lay.shapes, found)
            items.append((f"layout {mi}.{li}", found))
    return items

def moved_pairs(pairs):
    """Matched slides whose order changed. An insert or delete shifts indexes and is not a move."""
    matched = [(i, j) for i, j in pairs if i is not None and j is not None]
    rank_a = {p: r for r, p in enumerate(sorted(matched))}
    rank_b = {p: r for r, p in enumerate(matched)}
    return {p for p in matched if rank_a[p] != rank_b[p]}

def side_lines(src, mark):
    lines = []
    for sid, v in src[0].items():
        if v["txt"].strip():
            lines.append(f"  {mark} {sid} {v['name']} {v['txt']!r}")
        if v["strike"].strip():
            lines.append(f"  {mark} {sid} {v['name']} strike: {v['strike']!r}")
    if src[1].strip():
        lines.append(f"  {mark} notes: {src[1]!r}")
    lines += [f"  {mark} comment: {t!r}" for t in src[2]]
    return lines

def main():
    A, B = deck(sys.argv[1]), deck(sys.argv[2])
    print(f"slides {len(A)} -> {len(B)}")
    pairs = align(A, B)
    moved = moved_pairs(pairs)
    for i, j in pairs:
        if i is None or j is None:
            src = B[j] if i is None else A[i]
            print(f"[slide {(j if i is None else i) + 1}] only in {'B' if i is None else 'A'}")
            body = side_lines(src, "+" if i is None else "-")
            if body: print("\n".join(body))
            continue
        (a, an, ac), (b, bn, bc) = A[i], B[j]
        words, layout = [], []
        for sid in sorted(set(a) | set(b)):
            if sid not in b: words.append(f"  - removed {sid} {a[sid]['name']} {a[sid]['txt']!r}")
            elif sid not in a: words.append(f"  + added {sid} {b[sid]['name']} {b[sid]['geo']} fill={b[sid]['fill']} shadow={b[sid]['shadow']} {b[sid]['txt']!r}")
            else:
                for k in ("txt", "strike", "geo", "fill", "shadow", "img", "crop", "rot", "descr", "link"):
                    if a[sid][k] != b[sid][k]:
                        bucket = words if k in ("txt", "strike", "descr", "link") else layout
                        if k in ("txt", "strike"):
                            bucket += [f"  ~ {sid} {b[sid]['name']} {k}: {x!r} -> {y!r}" for x, y in excerpts(a[sid][k], b[sid][k])]
                        else:
                            bucket.append(f"  ~ {sid} {b[sid]['name']} {k}: {a[sid][k]!r} -> {b[sid][k]!r}")
                words += [f"  ~ {sid} {b[sid]['name']} mark: {x!r} -> {y!r}" for x, y in format_changes(a[sid]["styled"], b[sid]["styled"])]
        words += [f"  ~ notes: {x!r} -> {y!r}" for x, y in excerpts(an, bn)]
        for t, n in sorted((Counter(bc) - Counter(ac)).items()):
            words += [f"  + comment: {t!r}"] * n
        for t, n in sorted((Counter(ac) - Counter(bc)).items()):
            words += [f"  - comment: {t!r}"] * n
        lines = words + layout
        if lines or (i, j) in moved:
            print(f"[slide {j+1}]" + (f" (was {i+1})" if i != j else "") + (" moved" if (i, j) in moved else ""))
            if lines: print("\n".join(lines))
    # Masters and layouts are paired by name. Only their text can carry a comment.
    ca = dict(chrome(sys.argv[1]))
    cb = dict(chrome(sys.argv[2]))
    for name in list(ca) + [n for n in cb if n not in ca]:
        a, b = ca.get(name, {}), cb.get(name, {})
        lines = []
        for sid in sorted(set(a) | set(b), key=str):
            if sid not in a or sid not in b:
                src, sign = (b, "+") if sid in b else (a, "-")
                if src[sid]["txt"].strip():
                    lines.append(f"  {sign} {sid} {src[sid]['name']} {src[sid]['txt']!r}")
                continue
            for k in ("txt", "strike", "descr"):
                if a[sid][k] != b[sid][k]:
                    lines += [f"  ~ {sid} {b[sid]['name']} {k}: {x!r} -> {y!r}" for x, y in excerpts(a[sid][k], b[sid][k])]
            lines += [f"  ~ {sid} {b[sid]['name']} mark: {x!r} -> {y!r}" for x, y in format_changes(a[sid]["styled"], b[sid]["styled"])]
        if lines:
            print(f"[{name}]")
            print("\n".join(lines))

if __name__ == "__main__":
    main()
