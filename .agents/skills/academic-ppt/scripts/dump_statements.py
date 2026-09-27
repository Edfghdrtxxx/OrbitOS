#!/usr/bin/env python3
"""dump_statements.py <deck.pptx> <out.md>

Slide-by-slide dump of every statement for a precision review: text boxes, the LaTeX
source of each rendered equation/card PNG (via the deck folder's render_eq.py SPECS and
assets/eq/meta.json), figure file names, and speaker notes. Run from the deck folder.
PowerPoint re-encodes (and may crop) media on save: PNGs match by hash, then by pixel size;
anything unmatched is listed as [figure] and read from the slide image by the reviewer.
"""
import glob, hashlib, json, os, sys

from pptx import Presentation

sys.path.insert(0, os.getcwd())
import render_eq  # noqa: E402  (deck-local copy)

deck, out = sys.argv[1], sys.argv[2]
E = 914400
spec = {s["name"]: " ".join(s["body"].split()) for s in render_eq.SPECS}
meta = json.load(open("assets/eq/meta.json"))
by_hash = {hashlib.md5(open(f, "rb").read()).hexdigest(): os.path.basename(f)[:-4] for f in glob.glob("assets/eq/*.png")}
by_size = {}
for n, m in meta.items():
    by_size.setdefault((m["wPx"], m["hPx"]), []).append(n)
fig_hash = {hashlib.md5(open(f, "rb").read()).hexdigest(): os.path.basename(f) for f in glob.glob("assets/*.*")}

lines = []
for i, s in enumerate(Presentation(deck).slides, 1):
    lines.append(f"\n### Slide {i}\n")
    for sh in sorted(s.shapes, key=lambda x: (round(x.top / E, 1), x.left)):
        if sh.has_text_frame and sh.text_frame.text.strip():
            lines.append(f"- [text] {sh.text_frame.text.strip().replace(chr(10), ' / ')}")
        elif sh.shape_type == 13:
            h = hashlib.md5(sh.image.blob).hexdigest()
            name = by_hash.get(h)
            if not name:
                cand = by_size.get(sh.image.size, [])
                name = cand[0] if len(cand) == 1 else None
            if name and name in spec:
                lines.append(f"- [rendered LaTeX: {name}] {spec[name][:900]}")
            else:
                lines.append(f"- [figure: {fig_hash.get(h, 'embedded image')}]")
    if s.has_notes_slide and s.notes_slide.notes_text_frame is not None:
        lines.append(f"- [speaker notes] {s.notes_slide.notes_text_frame.text.strip()}")
open(out, "w").write("\n".join(lines))
print(f"{len(lines)} lines → {out}")
