#!/usr/bin/env python3
"""apply_chrome.py — run after `node build_dreb_pptx.js`.

pptxgenjs cannot write the presenter's gradient header rule, so the builder draws a placeholder
line named "HeaderRule". This swaps each one for the exact <p:cxnSp> from the presenter's v2 deck
(slide 3 = content slides, rule at y 528955 EMU; slide 2 = contents slide, y 699135 EMU).
The v2 deck is kept at assets/src/; DREB2026_SFS_last.pptx is the last delivered prototype.
"""
import copy
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
SRC = ROOT / "assets/src/DREB_Meeting_SFS_v2.pptx"
OUT = ROOT / "DREB2026_SFS.pptx"


def rule_from(slide):
    for el in slide.shapes._spTree.iter(qn("p:cxnSp")):
        return el
    raise SystemExit("no header rule in source slide")


src = Presentation(SRC)
rules = {int(rule_from(src.slides[i]).find(".//" + qn("a:off")).get("y")): rule_from(src.slides[i]) for i in (1, 2)}

deck = Presentation(OUT)
n = 0
for slide in deck.slides:
    tree = slide.shapes._spTree
    ids = [int(e.get("id")) for e in tree.iter(qn("p:cNvPr"))]
    for sp in list(tree.iter(qn("p:sp"))):
        nv = sp.find(".//" + qn("p:cNvPr"))
        if nv is None or nv.get("name") != "HeaderRule":
            continue
        y = int(sp.find(".//" + qn("a:off")).get("y"))
        key = min(rules, key=lambda k: abs(k - y))
        new = copy.deepcopy(rules[key])
        new.find(".//" + qn("p:cNvPr")).set("id", str(max(ids) + 1))
        ids.append(max(ids) + 1)
        sp.addprevious(new)
        tree.remove(sp)
        n += 1
deck.save(OUT)
print(f"{n} header rules replaced → {OUT.name}")
